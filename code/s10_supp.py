import os, json, pandas as pd
A=os.path.expanduser('~/rc_nafld_mr/analysis'); D=A+'/primary_clumped52'
LAB={'GCST90091033':'Ghodsian','GCST90054782':'Fairfield','finn-b-NAFLD':'FinnGen'}
out=A+'/Supplementary_Tables.xlsx'
w=pd.ExcelWriter(out,engine='openpyxl')
# S1 instruments
clump=pd.read_csv(os.path.expanduser('~/rc_nafld_mr/work/instr_clumped.tsv'),sep='\t')
base=clump[['rsid','chr','pos','p']].rename(columns={'rsid':'SNP','p':'P_exposure'})
for o,n in LAB.items():
    h=pd.read_csv(f'{D}/harmonised_{o}.tsv',sep='\t')
    base=base.merge(h[['SNP']].assign(**{f'used_{n}':'yes'}),on='SNP',how='left')
ref=pd.read_csv(f'{D}/harmonised_GCST90091033.tsv',sep='\t')
for o in list(LAB)[1:]:
    ref=pd.concat([ref,pd.read_csv(f'{D}/harmonised_{o}.tsv',sep='\t')]).drop_duplicates('SNP')
base=base.merge(ref[['SNP','ea','oa','eaf_x','bx','sex','F','r2_x']].rename(columns={'ea':'effect_allele','oa':'other_allele','eaf_x':'EAF','bx':'beta_RC_SD','sex':'SE_RC','F':'F_stat','r2_x':'R2'}),on='SNP',how='left')
base.fillna('no').to_excel(w,'S1 Instruments',index=False)
# S2 per-variant
for o,n in LAB.items():
    h=pd.read_csv(f'{D}/harmonised_{o}.tsv',sep='\t'); h.to_excel(w,f'S2 Variants {n}',index=False)
# S3 all MR estimates
rs=[]
for d,s in [(D,'Primary 52-variant set'),(A+'/sens_noAPOE','APOE region excluded'),(A+'/sens_original48','Original 48-variant set')]:
    t=pd.read_csv(d+'/mr_results.tsv',sep='\t'); t.insert(0,'instrument_set',s); t['outcome']=t.outcome.map(LAB); rs.append(t)
pd.concat(rs).to_excel(w,'S3 All MR estimates',index=False)
# S4 summaries
rows=[]
for d,s in [(D,'Primary 52-variant set'),(A+'/sens_noAPOE','APOE region excluded'),(A+'/sens_original48','Original 48-variant set')]:
    for x in json.load(open(d+'/summary.json')):
        r={k:v for k,v in x.items() if k not in('log','desc')}; r['outcome']=LAB[x['outcome']]; r['instrument_set']=s
        r['presso_outliers']=', '.join(r['presso_outliers']); r['steiger_removed']=', '.join(r['steiger_removed']); rows.append(r)
pd.DataFrame(rows).to_excel(w,'S4 Heterogeneity-pleiotropy',index=False)
# S5 LOO
for o,n in LAB.items():
    pd.read_csv(f'{D}/leave_one_out_{o}.tsv',sep='\t').to_excel(w,f'S5 LOO {n}',index=False)
# S6 proxies
pd.read_csv(os.path.expanduser('~/rc_nafld_mr/work/proxies_used.tsv'),sep='\t').to_excel(w,'S6 Proxies used',index=False)
w.close()
from openpyxl import load_workbook
wb=load_workbook(out)
for ws in wb:
    for c in ws[1]: c.font=c.font.copy(bold=True)
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=max(10,min(28,max(len(str(c.value)) if c.value is not None else 0 for c in col[:30])+2))
    ws.freeze_panes='A2'
wb.save(out); print(wb.sheetnames)
